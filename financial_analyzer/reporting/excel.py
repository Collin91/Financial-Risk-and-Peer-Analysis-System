"""Formatted Excel workbook: risk summary, peer comparison, trends, rule detail, events, data and methodology."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from financial_analyzer.analysis import risk
from financial_analyzer.analysis.metrics import METRICS_BY_KEY, displays_equal
from financial_analyzer.analysis.pipeline import PERIOD_DISCLOSURE
from financial_analyzer.data.standardize import LINE_ITEMS, SUPPORT_ITEMS, long_date, short_date
from financial_analyzer.events import DISCLAIMER as EVENTS_DISCLAIMER

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SUBHEADER_FILL = PatternFill("solid", fgColor="E8EDF3")
TITLE_FONT = Font(bold=True, size=14)
SECTION_FONT = Font(bold=True, size=12)
SUBTITLE_FONT = Font(italic=True, color="52514E")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
GOOD_FILL = PatternFill("solid", fgColor="E2F4E2")
BAD_FILL = PatternFill("solid", fgColor="FBE3E1")
WARN_FILL = PatternFill("solid", fgColor="FDF1D6")
LEVEL_FILLS = {
    "Lower Risk": PatternFill("solid", fgColor="C8EBC8"),
    "Moderate Risk": PatternFill("solid", fgColor="FDE7B0"),
    "Elevated Risk": PatternFill("solid", fgColor="F4C0BD"),
}
STATUS_FILLS = {risk.TRIGGERED: PatternFill("solid", fgColor="FBE3E1"),
                risk.INFORMATIONAL: PatternFill("solid", fgColor="EEF2F7")}

NUMBER_FORMATS = {"pct": "0.0%", "ratio": '0.00"x"', "days": '0" days"'}
# Differences: percentage metrics are shown in percentage points (value x 100).
DIFF_FORMATS = {"pct": '+0.0" pp";-0.0" pp"', "ratio": '+0.00"x";-0.00"x"', "days": '+0" days";-0" days"'}


def _header(ws: Worksheet, row: int, values: list, col: int = 1) -> None:
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=col + i, value=v)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _widths(ws: Worksheet, widths: dict[int, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def _write_frame(ws: Worksheet, df: pd.DataFrame, start_row: int = 1, formats: dict[str, str] | None = None) -> int:
    _header(ws, start_row, list(df.columns))
    for r, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=None if pd.isna(value) else value)
            cell.border = BORDER
            if formats and df.columns[c - 1] in formats:
                cell.number_format = formats[df.columns[c - 1]]
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(df.columns))}{start_row + len(df)}"
    return start_row + len(df) + 1


def _text(ws: Worksheet, row: int, text: str, font: Font | None = None, col: int = 1, merge_to: int = 0) -> int:
    cell = ws.cell(row=row, column=col, value=text)
    cell.alignment = WRAP
    if font:
        cell.font = font
    if merge_to:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=merge_to)
        ws.row_dimensions[row].height = max(15, 15 * (len(text) // 150 + 1))
    return row + 1


def _periods_table(ws: Worksheet, result, row: int, year: int) -> int:
    """The Company / Reported fiscal year / Period ended / Comparison year table."""
    _header(ws, row, ["Company", "Reported fiscal year", "Period ended", "Comparison year"])
    for name in result.company_names:
        p = result.period(name, year)
        if p is None:
            continue
        row += 1
        for c, v in enumerate([name, p.label, long_date(p.end), year], start=1):
            ws.cell(row=row, column=c, value=v).border = BORDER
    return row + 2


def _period_map(ws: Worksheet, result, row: int) -> int:
    """Company x comparison year grid of reported fiscal-year labels and period ends."""
    years = list(result.years)
    _header(ws, row, ["Reported period"] + [f"Comparison year {y}" for y in years])
    for name in result.company_names:
        row += 1
        ws.cell(row=row, column=1, value=name).font = BOLD
        for c, y in enumerate(years, start=2):
            p = result.period(name, y)
            ws.cell(row=row, column=c, value=f"{p.label} (YE {short_date(p.end)})" if p else "n/a").border = BORDER
    return row + 2


def _summary_sheet(wb: Workbook, result, notes: dict[str, list[str]]) -> None:
    ws = wb.active
    ws.title = "Summary"
    year = result.latest_year
    ws["A1"] = f"Financial Risk & Peer Analysis - {result.profile.name}"
    ws["A1"].font = TITLE_FONT
    target_label = result.period_label(result.target, year)
    row = _text(ws, 2, f"Company investigated: {result.target}   |   Peers: "
                       f"{', '.join(n for n in result.company_names if n != result.target)}   |   "
                       f"Comparison years {result.first_year}-{result.last_year}   |   Scores shown for comparison "
                       f"year {year} ({result.target} {target_label})", SUBTITLE_FONT)
    row = _text(ws, row, "Flags identify unusual financial patterns that may warrant further investigation. "
                         "They are not a fraud determination or a share-price forecast.", SUBTITLE_FONT)
    row = _text(ws, row + 1, PERIOD_DISCLOSURE, BOLD, merge_to=4)
    row = _periods_table(ws, result, row + 1, year)

    latest = result.scores[result.scores.comparison_year == year].set_index("company")
    ordered = sorted((n for n in result.company_names if n in latest.index), key=lambda n: -latest.loc[n, "score"])
    for name in ordered:
        score, level = int(latest.loc[name, "score"]), latest.loc[name, "risk_level"]
        label = (f"{name}{' (target)' if name == result.target else ''} - {result.period_label(name, year)}"
                 f" - {level}")
        ws.cell(row=row, column=1, value=label).font = SECTION_FONT
        ws.cell(row=row, column=1).fill = LEVEL_FILLS[level]
        ws.cell(row=row, column=2, value=f"Score: {score} points").font = BOLD
        ws.cell(row=row, column=2).fill = LEVEL_FILLS[level]
        row += 1
        fired = [r for r in result.rule_results if r.company == name and r.comparison_year == year and r.triggered]
        for r in sorted(fired, key=lambda r: -r.points):
            ws.cell(row=row, column=1, value=f"  • {r.detail}").alignment = WRAP
            ws.cell(row=row, column=2, value=f"+{r.points}  [{r.rule.id} {r.rule.name}] {r.rule.test}").font = SUBTITLE_FONT
            row += 1
        for note in notes.get(name, []):
            ws.cell(row=row, column=1, value=f"  • {note}").font = Font(color="1B6E1B")
            row += 1
        if not fired and not notes.get(name):
            ws.cell(row=row, column=1, value="  • No rules triggered.")
            row += 1
        row += 1

    if result.comparability:
        ws.cell(row=row, column=1, value="Keep in mind (full explanations on the Methodology sheet)").font = SECTION_FONT
        row += 1
        for w in result.comparability:
            cell = ws.cell(row=row, column=1, value=f"  • {w.summary or w.title}")
            cell.fill, cell.alignment = WARN_FILL, WRAP
            row += 1
        row += 1

    ws.cell(row=row, column=1, value="Score history by comparison year").font = SECTION_FONT
    row += 1
    years = list(result.years)
    _header(ws, row, ["Company"] + [f"Comparison year {y}" for y in years])
    history = result.scores.set_index(["company", "comparison_year"])
    for name in result.company_names:
        row += 1
        ws.cell(row=row, column=1, value=name)
        for c, y in enumerate(years, start=2):
            if (name, y) not in history.index:
                ws.cell(row=row, column=c, value="n/a")
                continue
            v = int(history.loc[(name, y), "score"])
            cell = ws.cell(row=row, column=c, value=v)
            cell.fill = LEVEL_FILLS[risk.risk_level(v)]
            cell.alignment = Alignment(horizontal="center")
    ws.cell(row=row + 2, column=1, value="Bands: 0-2 Lower Risk, 3-5 Moderate Risk, 6+ Elevated Risk. See 'Rule "
                                         "Details' for every test and 'Methodology' for definitions.").font = SUBTITLE_FONT
    _widths(ws, {1: 78, 2: 70, 3: 24, 4: 24, 5: 24, 6: 24})


def _peer_sheet(wb: Workbook, result) -> None:
    ws = wb.create_sheet("Peer Comparison")
    year = result.latest_year
    table = result.peer_comparison(year)
    ws["A1"] = f"Peer comparison - comparison year {year}"
    ws["A1"].font = TITLE_FONT
    _text(ws, 2, PERIOD_DISCLOSURE, SUBTITLE_FONT, merge_to=8)
    ws["A3"] = (f"Green / red shading marks whether {result.target} is better or worse than the median of its peers. "
                "No shading where values match at report precision or the figures are not comparable (†).")
    ws["A3"].font = SUBTITLE_FONT
    cols = ["Metric"] + list(table.columns) + [f"{result.target} vs median", "Comparability", "Formula"]
    _header(ws, 5, cols)
    for c, col in enumerate(table.columns, start=2):  # period row under the company names
        p = result.period(col, year)
        cell = ws.cell(row=6, column=c, value=f"{p.label} · YE {short_date(p.end)}" if p else "")
        cell.fill, cell.font = SUBHEADER_FILL, Font(italic=True, size=9)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    r = 6
    target_col = 2 + list(table.columns).index(result.target)
    for key, values in table.iterrows():
        r += 1
        metric = METRICS_BY_KEY[key]
        blocked = result.blocked(key, result.company_names)
        ws.cell(row=r, column=1, value=metric.label + (" †" if blocked else "")).font = BOLD
        for c, v in enumerate(values, start=2):
            ws.cell(row=r, column=c, value=None if pd.isna(v) else float(v)).number_format = NUMBER_FORMATS[metric.fmt]
        tv, med = values.get(result.target), values["Peer median (excl. target)"]
        if not (pd.isna(tv) or pd.isna(med)):
            diff = ws.cell(row=r, column=len(cols) - 2, value=float((tv - med) * (100 if metric.fmt == "pct" else 1)))
            diff.number_format = DIFF_FORMATS[metric.fmt]
        marker = result.peer_marker(key, year)
        if marker:
            ws.cell(row=r, column=target_col).fill = GOOD_FILL if marker == "better" else BAD_FILL
        if blocked:
            ws.cell(row=r, column=len(cols) - 1, value="Not directly comparable - informational").fill = WARN_FILL
        ws.cell(row=r, column=len(cols), value=metric.formula).font = SUBTITLE_FONT
        for c in range(1, len(cols) + 1):
            ws.cell(row=r, column=c).border = BORDER

    r += 2
    for w in result.comparability:
        if w.suppresses_peer_points:
            ws.cell(row=r, column=1, value=f"† {w.title}").font = BOLD
            ws.cell(row=r, column=2, value=w.message).alignment = WRAP
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=len(cols))
            ws.row_dimensions[r].height = 15 * (len(w.message) // 180 + 1)
            r += 1

    r += 2
    first = result.period(result.target, result.first_year)
    prior = result.period(result.target, year - 1)
    latest_label = result.period_label(result.target, year)
    span = f"{first.label}-{prior.label}" if first and prior else "prior periods"
    ws.cell(row=r, column=1, value=f"{result.target} vs its own history ({latest_label} vs {span})").font = SECTION_FONT
    r += 1
    hist = result.target_vs_history()
    _header(ws, r, ["Metric", result.period_label(result.target, year, with_date=False), "Prior average", "Prior low",
                    "Prior high", "Change vs average"])
    for key, values in hist.iterrows():
        r += 1
        metric = METRICS_BY_KEY[key]
        fmt = NUMBER_FORMATS[metric.fmt]
        ws.cell(row=r, column=1, value=metric.label).font = BOLD
        for c, col in enumerate(["latest", "prior_average", "prior_min", "prior_max"], start=2):
            v = values[col]
            ws.cell(row=r, column=c, value=None if pd.isna(v) else float(v)).number_format = fmt
        if not (pd.isna(values["latest"]) or pd.isna(values["prior_average"])):
            change = values["latest"] - values["prior_average"]
            cell = ws.cell(row=r, column=6, value=float(change * 100 if metric.fmt == "pct" else change))
            cell.number_format = DIFF_FORMATS[metric.fmt]
            if metric.higher_is_better is not None and not displays_equal(values["latest"], values["prior_average"],
                                                                           metric.fmt):
                cell.fill = GOOD_FILL if (change > 0) == metric.higher_is_better else BAD_FILL
        for c in range(1, 7):
            ws.cell(row=r, column=c).border = BORDER
    _widths(ws, {1: 34, **{i: 18 for i in range(2, len(cols) - 1)}, len(cols) - 1: 26, len(cols): 60})


def _trends_sheet(wb: Workbook, result) -> None:
    ws = wb.create_sheet("Trends")
    ws["A1"] = "Metric trends by company and comparison year"
    ws["A1"].font = TITLE_FONT
    _text(ws, 2, PERIOD_DISCLOSURE, SUBTITLE_FONT, merge_to=7)
    r = _period_map(ws, result, 4)
    years = list(result.years)
    for key in result.profile.metric_keys:
        metric = METRICS_BY_KEY[key]
        blocked = result.blocked(key, result.company_names)
        ws.cell(row=r, column=1, value=metric.label + (" †" if blocked else "")).font = SECTION_FONT
        ws.cell(row=r, column=2, value=metric.formula).font = SUBTITLE_FONT
        r += 1
        _header(ws, r, ["Company"] + [f"Comparison year {y}" for y in years])
        for name in result.company_names:
            r += 1
            ws.cell(row=r, column=1, value=name).font = BOLD if name == result.target else Font()
            series = result.metrics.xs(name, level="company")[key]
            for c, y in enumerate(years, start=2):
                v = series.get(y)
                cell = ws.cell(row=r, column=c, value=None if v is None or pd.isna(v) else float(v))
                cell.number_format = NUMBER_FORMATS[metric.fmt]
                cell.border = BORDER
        r += 2
    _widths(ws, {1: 22, **{i: 24 for i in range(2, len(years) + 2)}})


def _rules_sheet(wb: Workbook, result) -> None:
    ws = wb.create_sheet("Rule Details")
    df = risk.results_frame(result.rule_results).sort_values(["comparison_year", "company", "rule_id"],
                                                               ascending=[False, True, True])
    ws["A1"] = "Every rule evaluated, for every company and period"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = ("Filter 'result': TRIGGERED added points; 'informational' peer comparisons assign no points; "
                "'not evaluated' means data was missing or the test did not apply.")
    ws["A2"].font = SUBTITLE_FONT
    end = _write_frame(ws, df, start_row=4)
    status_col = list(df.columns).index("result") + 1
    for r in range(5, end):
        fill = STATUS_FILLS.get(ws.cell(row=r, column=status_col).value)
        if fill:
            for c in range(1, len(df.columns) + 1):
                ws.cell(row=r, column=c).fill = fill
    _widths(ws, {1: 10, 2: 11, 3: 12, 4: 11, 5: 8, 6: 30, 7: 16, 8: 50, 9: 13, 10: 7, 11: 10, 12: 90, 13: 40})


def _events_sheet(wb: Workbook, result) -> None:
    ws = wb.create_sheet("Events")
    ws["A1"] = "Potential Explanatory Events"
    ws["A1"].font = TITLE_FONT
    row = _text(ws, 2, EVENTS_DISCLAIMER, BOLD, merge_to=8)
    if not result.events_run:
        _text(ws, row + 1, "Event retrieval was not run for this report.", SUBTITLE_FONT)
        return
    for err in result.event_errors:
        row = _text(ws, row, f"Retrieval note: {err}", SUBTITLE_FONT)
    row += 1
    ws.cell(row=row, column=1, value="Significant changes detected").font = SECTION_FONT
    row += 1
    changes = pd.DataFrame([vars(c) for c in result.event_changes])
    if not changes.empty:
        row = _write_frame(ws, changes, start_row=row) + 1
    ws.cell(row=row, column=1, value="Candidate events (up to three per company and period)").font = SECTION_FONT
    row += 1
    events = pd.DataFrame([{
        "company": e.company, "reported_fiscal_year": e.reported_fiscal_year, "comparison_year": e.comparison_year,
        "event_date": e.event_date, "filing_type": e.filing_type, "relevance": e.relevance,
        "affected_metric": e.affected_metric, "description": e.description, "sec_source": e.url,
    } for e in result.events])
    if events.empty:
        _text(ws, row, "No candidate events matched.", SUBTITLE_FONT)
        return
    start = row
    end = _write_frame(ws, events, start_row=row)
    for r in range(start + 1, end):
        cell = ws.cell(row=r, column=9)
        cell.hyperlink, cell.font = cell.value, Font(color="1F5FBF", underline="single")
        ws.cell(row=r, column=8).alignment = WRAP
    ws.freeze_panes = None
    _widths(ws, {1: 12, 2: 12, 3: 11, 4: 12, 5: 10, 6: 10, 7: 26, 8: 90, 9: 60})


def _methodology_sheet(wb: Workbook, result) -> None:
    ws = wb.create_sheet("Methodology")
    ws["A1"] = "Methodology"
    ws["A1"].font = TITLE_FONT
    lines = [
        ("Data source", "Annual reports (10-K for US filers, 20-F for foreign filers) downloaded from SEC EDGAR. "
                        "Values are read from each filing's XBRL instance; only consolidated (non-segment) facts are "
                        "used. When a later filing restates a prior period, the most recently filed value is used."),
        ("Standardization", "Each line item has an ordered list of candidate XBRL concepts across US GAAP and IFRS. "
                            "The first concept reported for the period is used; documented derivations fill gaps. "
                            "The 'Data Lineage' sheet shows the concept, filing and FX rate behind every number."),
        ("Fiscal years", "Each company's reported fiscal-year label is taken from its own annual report "
                         "(dei:DocumentFiscalYearFocus) and never changed. " + PERIOD_DISCLOSURE +
                         " The comparison year is the calendar year containing most of the fiscal period; for "
                         "example Toyota FY2026 (year ended March 31, 2026) has comparison year 2025."),
        ("Currency", "Ratios and growth rates are calculated in each company's reporting currency, so exchange-rate "
                     "moves do not distort them. USD amounts use Federal Reserve H.10 rates from FRED: period-average "
                     "rates for income and cash-flow items, period-end rates for balance-sheet items."),
        ("Capital expenditures", "Cash paid for purchases of property, plant and equipment, for every company. "
                                 "Acquisitions, finance receivables, investments, intangible assets and vehicles bought "
                                 "for leasing to customers are excluded. A broader 'productive assets' concept is "
                                 "accepted only when the filer's intangible assets are immaterial (1% of total assets "
                                 "or less); otherwise capex is left missing and capex rules are not evaluated."),
        ("Peer median", f"For each company, the median of the other companies' values for the same comparison year. "
                        f"Peer-based rules assign points only with at least {risk.MIN_PEERS} usable, comparable peer "
                        f"values; otherwise the comparison is shown as informational."),
        ("Risk score", "Sum of the points for every triggered rule. 0-2 Lower Risk, 3-5 Moderate Risk, 6+ Elevated "
                       "Risk."),
        ("Limitations", "The analysis flags unusual financial patterns for follow-up. It does not detect fraud, "
                        "assess management intent, or predict share prices. Accounting-policy differences between "
                        "US GAAP and IFRS remain even after label matching."),
    ]
    r = 3
    for title, text in lines:
        ws.cell(row=r, column=1, value=title).font = BOLD
        ws.cell(row=r, column=2, value=text).alignment = WRAP
        r += 1
    for note in result.profile.notes:
        ws.cell(row=r, column=1, value=f"{result.profile.name} note").font = BOLD
        ws.cell(row=r, column=2, value=note).alignment = WRAP
        r += 1
    for w in result.comparability:
        ws.cell(row=r, column=1, value=f"Comparability: {w.title}").font = BOLD
        ws.cell(row=r, column=1).alignment = WRAP
        ws.cell(row=r, column=2, value=w.message).alignment = WRAP
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Rule families").font = SECTION_FONT
    for prefix, meaning in risk.RULE_FAMILIES.items():
        r += 1
        ws.cell(row=r, column=1, value=prefix).font = BOLD
        ws.cell(row=r, column=2, value=meaning)

    r += 2
    ws.cell(row=r, column=1, value="Rule catalogue").font = SECTION_FONT
    r += 1
    _header(ws, r, ["Rule", "Name", "Points", "Test"])
    for rule in result.profile.rules:
        r += 1
        for c, v in enumerate([rule.id, rule.name, rule.points, rule.test], start=1):
            ws.cell(row=r, column=c, value=v).alignment = WRAP

    r += 2
    ws.cell(row=r, column=1, value="Metrics").font = SECTION_FONT
    r += 1
    _header(ws, r, ["Metric", "Formula"])
    for key in result.profile.metric_keys:
        r += 1
        ws.cell(row=r, column=1, value=METRICS_BY_KEY[key].label)
        ws.cell(row=r, column=2, value=METRICS_BY_KEY[key].formula)

    r += 2
    ws.cell(row=r, column=1, value="Line items and candidate XBRL concepts").font = SECTION_FONT
    r += 1
    _header(ws, r, ["Line item", "Candidates (priority order)"])
    for item in LINE_ITEMS + SUPPORT_ITEMS:
        r += 1
        ws.cell(row=r, column=1, value=item.label).alignment = WRAP
        ws.cell(row=r, column=2, value=", ".join(item.candidates)).alignment = WRAP

    if result.warnings:
        r += 2
        ws.cell(row=r, column=1, value="Data-quality notes").font = SECTION_FONT
        for w in result.warnings:
            r += 1
            ws.cell(row=r, column=2, value=w).alignment = WRAP
    r += 2
    ws.cell(row=r, column=1, value="Periods and fiscal-year label sources").font = SECTION_FONT
    r += 1
    _header(ws, r, ["Company / fiscal year", "Period", "Comparison year", "Label source"])
    for p in result.periods.itertuples(index=False):
        r += 1
        base = " (base year)" if p.is_base_year else ""
        for c, v in enumerate([f"{p.company} {p.reported_fiscal_year}{base}", f"{p.period_start} to {p.period_end}",
                               p.comparison_year, p.fiscal_year_label_source], start=1):
            ws.cell(row=r, column=c, value=v)
    _widths(ws, {1: 30, 2: 110, 3: 8, 4: 80})


def write_workbook(result, path: Path, notes: dict[str, list[str]]) -> Path:
    """Sheets: Summary, Peer Comparison, Trends, Rule Details, Events (if run), Data, Data Lineage, Methodology."""
    wb = Workbook()
    _summary_sheet(wb, result, notes)
    _peer_sheet(wb, result)
    _trends_sheet(wb, result)
    _rules_sheet(wb, result)
    if result.events_run:
        _events_sheet(wb, result)

    ws = wb.create_sheet("Data")
    money = {c: '#,##0' for c in result.cleaned.columns if c.endswith("_usd_m")}
    money.update({"fx_avg_rate_per_usd": "0.00", "fx_period_end_rate_per_usd": "0.00"})
    _write_frame(ws, result.cleaned, formats=money)
    _widths(ws, {i: 15 for i in range(1, len(result.cleaned.columns) + 1)})

    ws = wb.create_sheet("Data Lineage")
    _write_frame(ws, result.lineage, formats={"native_value_m": "#,##0", "usd_value_m": "#,##0", "fx_rate_used": "0.00"})
    _widths(ws, {1: 10, 2: 11, 3: 12, 4: 12, 5: 11, 6: 30, 7: 11, 8: 15, 9: 9, 10: 11, 11: 14, 12: 14, 13: 70,
                 14: 60, 15: 40})

    _methodology_sheet(wb, result)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
